#!/usr/bin/env python3
"""Render a deck and find what a reader would call clumsy.

    python3 tools/deckcheck.py deck.pptx

Renders through LibreOffice, reads every word's box from the PDF with
pdftotext, and reports three kinds of fault per slide:

- collision: two words drawn on top of each other, which is how the matrix
  legend landed on table rows and a long question landed on its own label;
- off slide: a word below the body area, which is text running off the page;
- escape: a word that starts inside a card and ends below that card's edge.

The measurement the builder uses is an estimate. This check reads what was
actually rendered, so it is the arbiter when the two disagree.
"""

from __future__ import annotations

import os
import re
import shutil
import subprocess
import sys
import tempfile
from html import unescape

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU_PER_PT = 12700
# The footer's two small lines start at 5.25in. Anything above that line is body
# copy, and body copy reaching into it has run off the slide.
FOOTER_TOP_PT = 5.22 * 72


def render(path, workdir):
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir',
                    workdir, path], check=True, capture_output=True, timeout=240)
    pdf = os.path.join(workdir, os.path.splitext(os.path.basename(path))[0] + '.pdf')
    html = os.path.join(workdir, 'words.html')
    subprocess.run(['pdftotext', '-bbox', pdf, html], check=True, capture_output=True)
    return html


_PAGE = re.compile(r'<page width="([\d.]+)" height="([\d.]+)">(.*?)</page>', re.S)
_WORD = re.compile(r'<word xMin="([\d.]+)" yMin="([\d.]+)" xMax="([\d.]+)" '
                   r'yMax="([\d.]+)">(.*?)</word>')


def words_by_page(html):
    text = open(html, encoding='utf8').read()
    pages = []
    for width, height, body in _PAGE.findall(text):
        words = [(float(a), float(b), float(c), float(d), unescape(w))
                 for a, b, c, d, w in _WORD.findall(body)]
        pages.append((float(width), float(height), words))
    return pages


def panels(slide, scale):
    """Card rectangles, in rendered points."""
    out = []
    for shape in slide.shapes:
        if shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE:
            continue
        w, h = shape.width / EMU_PER_PT * scale, shape.height / EMU_PER_PT * scale
        # Cards only: skip slivers, pills and the full slide background.
        if w < 60 or h < 40 or (w > 700 * scale and h > 390 * scale):
            continue
        x, y = shape.left / EMU_PER_PT * scale, shape.top / EMU_PER_PT * scale
        out.append((x, y, x + w, y + h))
    return out


def overlap(a, b):
    ix = min(a[2], b[2]) - max(a[0], b[0])
    iy = min(a[3], b[3]) - max(a[1], b[1])
    if ix <= 0 or iy <= 0:
        return 0.0
    smaller = min((a[2] - a[0]) * (a[3] - a[1]), (b[2] - b[0]) * (b[3] - b[1]))
    return (ix * iy) / smaller if smaller else 0.0


def check(path):
    deck = Presentation(path)
    workdir = tempfile.mkdtemp()
    try:
        pages = words_by_page(render(path, workdir))
    finally:
        shutil.rmtree(workdir, ignore_errors=True)
    faults = []
    for number, ((width, height, words), slide) in enumerate(zip(pages, deck.slides), 1):
        scale = width / (deck.slide_width / EMU_PER_PT)
        footer_top = FOOTER_TOP_PT * scale
        cards = panels(slide, scale)
        # The footer is the only text allowed in the footer zone.
        body = [w for w in words if w[1] < footer_top]
        for i, a in enumerate(body):
            for b in body[i + 1:]:
                if overlap(a, b) > 0.25:
                    faults.append((number, 'collision', '%r over %r' % (a[4], b[4])))
        for word in body:
            if word[3] > footer_top:
                faults.append((number, 'off slide', repr(word[4])))
            for x0, y0, x1, y1 in cards:
                starts_inside = x0 <= word[0] <= x1 and y0 <= word[1] <= y1
                if starts_inside and word[3] > y1 + 1.5:
                    faults.append((number, 'escape', '%r leaves its card' % word[4]))
                    break
    return len(deck.slides), faults


def main(argv):
    if not argv:
        print(__doc__)
        return 2
    slides, faults = check(argv[0])
    by_slide = {}
    for number, kind, detail in faults:
        by_slide.setdefault(number, []).append((kind, detail))
    print('%d slides, %d faults on %d slides' % (slides, len(faults), len(by_slide)))
    for number in sorted(by_slide):
        kinds = {}
        for kind, detail in by_slide[number]:
            kinds.setdefault(kind, []).append(detail)
        summary = ', '.join('%d %s' % (len(v), k) for k, v in kinds.items())
        example = next(iter(kinds.values()))[0]
        print('  slide %-3d %-32s e.g. %s' % (number, summary, example[:70]))
    return 1 if faults else 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1:]))
