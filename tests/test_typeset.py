"""Tests for the typesetting layer and for decks built from long content.

Run with: python3 -m pytest tests -q
"""

import os
import shutil
import sys

import pytest
from pptx import Presentation
from pptx.util import Inches

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, 'tools'))

from battlecards import library, service, typeset  # noqa: E402

WIDTH = int(Inches(3))
# The footer's small print starts below the body area, at about 5.19in.
FOOTER_TOP = Inches(5.1)


# ── sources out of body copy ────────────────────────────────────────────────

def test_a_source_sentence_moves_to_the_citations():
    clean, found = typeset.split_sources(
        'They plan by cluster. Source: https://example.com/page. Ask for a demo.')
    assert clean == 'They plan by cluster. Ask for a demo.'
    assert found == ['https://example.com/page']


def test_a_page_reference_keeps_the_full_stop():
    clean, found = typeset.split_sources(
        'They forecast rate of sale, per Planning_Guide.pdf pages 15 and 16. '
        'Do not call this a gap.')
    assert clean == 'They forecast rate of sale. Do not call this a gap.'
    assert found == ['Planning_Guide.pdf pages 15 and 16']


def test_ordinary_per_phrases_are_left_alone():
    text = 'Plans run per store and per SKU.'
    assert typeset.split_sources(text) == (text, [])


def test_citation_labels_are_short():
    labels = typeset.cites('https://www.example.com/a/very/long/path?x=1; '
                           'Planning_Guide.pdf pages 4, 5 and 6')
    assert labels[0] == ('example.com', 'https://www.example.com/a/very/long/path?x=1')
    assert labels[1] == ('Planning Guide, pp. 4 to 6', None)


def test_a_url_never_swallows_the_sentence_after_it():
    labels = typeset.cites('https://example.com/news Reach figures from Notes.docx, '
                           'section 1.6')
    urls = [url for _, url in labels if url]
    assert urls == ['https://example.com/news']


# ── fitting ─────────────────────────────────────────────────────────────────

def test_measure_counts_the_line_spacing():
    one = typeset.measure('Short line.', WIDTH, 10, spacing=1.0)
    loose = typeset.measure('Short line.', WIDTH, 10, spacing=1.5)
    assert loose == pytest.approx(one * 1.5, rel=0.01)


@pytest.mark.parametrize('repeat', [1, 4, 12, 40])
def test_fitted_prose_never_exceeds_its_box(repeat):
    text = ' '.join(['Planners set option counts by cluster and channel.'] * repeat)
    height = int(Inches(1.2))
    pt, shown, trimmed = typeset.fit_prose(text, WIDTH, height, 11, 8.5)
    assert 8.5 <= pt <= 11
    assert typeset.measure(shown, WIDTH, pt) <= height
    assert trimmed == (shown != text)


def test_trimming_keeps_whole_sentences():
    text = ' '.join('Sentence number %d is here.' % n for n in range(40))
    _, shown, trimmed = typeset.fit_prose(text, WIDTH, int(Inches(0.6)), 10, 9)
    assert trimmed
    assert shown.endswith('.') and not shown.endswith(typeset.ELLIPSIS)


def test_trimmed_text_grows_back_into_its_box():
    text = ' '.join('Sentence number %d is here.' % n for n in range(40))
    pt, _, _ = typeset.fit_prose(text, WIDTH, int(Inches(0.6)), 11, 8.5)
    assert pt > 8.5


def test_balance_leaves_no_orphan():
    assert typeset.balance(10, 3) == [3, 3, 2, 2]
    assert typeset.balance(4, 2) == [2, 2]
    assert typeset.balance(0, 3) == []


# ── a deck from long content ────────────────────────────────────────────────

LONG = ('Planners choose the clustering algorithm and the store attributes it '
        'uses, then set aggregation and granularity for each category. '
        'The buyer will ask how much setup each approach needs before the '
        'first plan is usable, so answer with the steps, not the adjectives. '
        'Source: https://www.example.com/assortment-planning. '
        'Both vendors describe localisation, per Planning_Guide.pdf pages 6 and 20.')


def _long_card():
    """Synthetic content as long as a generated card writes, never a real one."""
    card = library.scaffold('Acme Planning', 'AssortSmart')
    card['snapshot'].update({key: LONG for key in (
        'headquarters', 'founded', 'employees', 'ownership', 'funding',
        'target_segment', 'go_to_market', 'deployment')})
    card['snapshot']['recent_moves'] = [LONG] * 6
    card['positioning'] = {'their_claim': LONG * 2, 'our_claim': LONG, 'wedge': LONG}
    card['their_strengths'] = [LONG] * 9
    card['their_weaknesses'] = [LONG] * 7
    card['our_advantages'] = [{'title': 'An advantage with a long title that wraps',
                               'detail': LONG * 2,
                               'proof': 'https://www.example.com/a; Planning_Guide.pdf page 4'}
                              for _ in range(10)]
    card['comparison'] = [{'capability': 'Capability %d with a longer name' % n,
                           'ia': 'strong', 'competitor': 'partial', 'note': LONG}
                          for n in range(18)]
    card['objections'] = [{'objection': 'A long objection the buyer raises in the '
                                        'first meeting, word for word?',
                           'response': LONG * 2, 'proof': 'https://www.example.com/b'}
                          for _ in range(7)]
    card['landmines'] = [{'question': 'Ask them to show the calculation live, and '
                                      'where each input comes from, on one screen?',
                          'why': LONG, 'listen_for': LONG} for _ in range(8)]
    card['proof_points'] = [{'stat': '11 March 2026', 'label': 'A dated launch',
                             'detail': LONG, 'source': 'https://www.example.com/c'},
                            {'stat': '35%', 'label': 'A result',
                             'detail': LONG, 'source': 'https://www.example.com/d'}]
    card['talk_track'] = {'positioning': LONG * 2, 'elevator': LONG,
                          'discovery_open': LONG, 'trap': LONG * 3}
    card['dos'] = [LONG] * 8
    card['donts'] = [LONG] * 9
    card['next_steps'] = [LONG] * 6
    return card


@pytest.fixture(scope='module')
def long_deck(tmp_path_factory):
    out = tmp_path_factory.mktemp('long')
    result = service.build(_long_card(), str(out))
    return result['path']


def test_long_content_stays_inside_the_slide(long_deck):
    deck = Presentation(long_deck)
    limit = Inches(5.2)
    for number, slide in enumerate(deck.slides, 1):
        for shape in slide.shapes:
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            if shape.top >= FOOTER_TOP:              # the footer itself
                continue
            assert shape.top + shape.height <= limit, (
                'slide %d: %r runs past the body' % (number, shape.text_frame.text[:40]))


def test_long_content_keeps_a_legible_size(long_deck):
    deck = Presentation(long_deck)
    for number, slide in enumerate(deck.slides, 1):
        for shape in slide.shapes:
            frames = [shape.text_frame] if shape.has_text_frame else []
            if getattr(shape, 'has_table', False) and shape.has_table:
                frames = [cell.text_frame for row in shape.table.rows for cell in row.cells]
            for frame in frames:
                for para in frame.paragraphs:
                    for run in para.runs:
                        if run.font.size and run.text.strip() and shape.top < FOOTER_TOP:
                            assert run.font.size.pt >= 6.5, (
                                'slide %d: %r at %.1fpt' % (number, run.text[:30],
                                                            run.font.size.pt))


def test_trimmed_detail_is_kept_in_the_notes(long_deck):
    deck = Presentation(long_deck)
    notes = '\n'.join(slide.notes_slide.notes_text_frame.text for slide in deck.slides
                      if slide.has_notes_slide)
    assert 'IN FULL' in notes
    assert 'SOURCES' in notes
    assert 'example.com/assortment-planning' in notes


def test_citations_leave_the_slide_face(long_deck):
    deck = Presentation(long_deck)
    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                assert 'Source:' not in shape.text_frame.text
                assert 'Planning_Guide.pdf' not in shape.text_frame.text


def test_long_sections_paginate_with_a_counter(long_deck):
    deck = Presentation(long_deck)
    texts = [' '.join(shape.text_frame.text for shape in slide.shapes
                      if shape.has_text_frame) for slide in deck.slides]
    head_to_head = [t for t in texts if 'Head to head' in t]
    assert len(head_to_head) >= 2
    assert all(' of %d' % len(head_to_head) in t.lower() for t in head_to_head)


@pytest.mark.skipif(not (shutil.which('soffice') and shutil.which('pdftotext')),
                    reason='needs LibreOffice and poppler to render')
def test_rendered_long_deck_has_no_collisions_or_overflow(long_deck):
    import deckcheck
    _, faults = deckcheck.check(long_deck)
    assert not faults, faults[:5]


# ── the comparison matrix ───────────────────────────────────────────────────

def test_the_edge_says_who_leads_a_row():
    from battlecards.builder import BattlecardDeck
    edge = BattlecardDeck.edge
    assert edge({'ia': 'strong', 'competitor': 'partial'}) == 'ours'
    assert edge({'ia': 'partial', 'competitor': 'partial'}) == 'level'
    assert edge({'ia': 'none', 'competitor': 'strong'}) == 'theirs'
    assert edge({'ia': 'strong', 'competitor': 'unknown'}) == 'verify'


def _matrix_slides(path):
    deck = Presentation(path)
    return [slide for slide in deck.slides
            if any(shape.has_text_frame and 'Head to head' in shape.text_frame.text
                   for shape in slide.shapes)]


def test_the_matrix_rates_with_harvey_balls_not_words(tmp_path):
    card = library.scaffold('Acme Planning', 'AssortSmart')
    card['comparison'] = [
        {'capability': 'Cluster plans', 'ia': 'strong', 'competitor': 'partial', 'note': 'x'},
        {'capability': 'Pack sizing', 'ia': 'partial', 'competitor': 'none', 'note': 'y'},
        {'capability': 'Open to buy', 'ia': 'strong', 'competitor': 'unknown', 'note': 'z'},
    ]
    slides = _matrix_slides(service.build(card, str(tmp_path))['path'])
    assert slides
    names = [shape.name for slide in slides for shape in slide.shapes]
    for rating in ('Strong', 'Partial', 'Gap', 'Unclear'):
        assert 'Rating: %s' % rating in names
    assert any(name == 'Rating: Partial, fill' for name in names)
    words = {shape.text_frame.text.strip() for slide in slides for shape in slide.shapes
             if shape.has_text_frame}
    assert not words & {'Strong', 'Partial', 'Gap', 'Unclear'}


def test_the_matrix_leads_with_our_advantages(tmp_path):
    card = library.scaffold('Acme Planning', 'AssortSmart')
    card['comparison'] = [
        {'capability': 'Unknown row', 'ia': 'unknown', 'competitor': 'strong', 'note': ''},
        {'capability': 'Their row', 'ia': 'partial', 'competitor': 'strong', 'note': ''},
        {'capability': 'Our row', 'ia': 'strong', 'competitor': 'none', 'note': ''},
    ]
    slide = _matrix_slides(service.build(card, str(tmp_path))['path'])[0]
    edges = [shape.name[len('Edge: '):] for shape in slide.shapes
             if shape.name.startswith('Edge: ')]
    assert edges == ['ours', 'theirs', 'verify']


# ── the head to head deck ───────────────────────────────────────────────────

def _h2h_card(rows):
    card = library.scaffold('Acme Planning', 'AssortSmart')
    card['comparison'] = rows
    card['meta']['depth'] = 'matrix'
    return card


H2H_ROWS = [
    {'capability': 'Cluster plans', 'ia': 'strong', 'competitor': 'partial',
     'note': 'They cluster by region. Ask to see a store level plan.'},
    {'capability': 'Pack sizing', 'ia': 'partial', 'competitor': 'partial',
     'note': 'Both ship it. Compete on setup time.'},
    {'capability': 'Open to buy', 'ia': 'none', 'competitor': 'strong',
     'note': 'They lead here. Move the buyer to the assortment outcome.'},
    {'capability': 'Image generation', 'ia': 'strong', 'competitor': 'unknown',
     'note': 'Nothing found either way. Confirm before the demo.'},
]


def test_head_to_head_is_its_own_short_deck(tmp_path):
    result = service.build(_h2h_card(H2H_ROWS), str(tmp_path))
    assert result['filename'].startswith('IA_Head_to_Head_Acme_Planning')
    assert result['card']['options']['sections'] == [
        'matrix_cover', 'matrix_map', 'comparison', 'matrix_plays']
    assert 4 <= result['slide_count'] <= 10
    assert result['compatibility']['ok']
    deck = Presentation(result['path'])
    text = ' '.join(shape.text_frame.text for slide in deck.slides
                    for shape in slide.shapes if shape.has_text_frame)
    assert 'Competitive battlecard'.upper() not in text, 'no second cover'
    assert 'AssortSmart leads on 1 of 4 capabilities' in text
    assert 'at a glance' in text and 'play the matrix' in text


def test_the_scoreboard_marks_every_capability(tmp_path):
    result = service.build(_h2h_card(H2H_ROWS), str(tmp_path))
    cover = Presentation(result['path']).slides[0]
    marks = [shape.name for shape in cover.shapes if shape.name.startswith('Capability: ')]
    assert len(marks) == len(H2H_ROWS)
    edges = [mark.rsplit('(', 1)[1].rstrip(')') for mark in marks]
    assert edges == ['ours', 'level', 'theirs', 'verify']


def test_the_play_is_the_last_instruction_in_the_note():
    from battlecards.builder import BattlecardDeck
    deck = BattlecardDeck(_h2h_card(H2H_ROWS))
    lines = deck._play_lines('Overlap is real. Do not claim absence. Compete on how it '
                             'is produced.')
    assert lines[0] == 'Compete on how it is produced.'
    assert deck._play_lines('A plain description.') == ['A plain description.']


def test_other_lengths_never_carry_the_head_to_head_slides():
    from battlecards import schema
    for key, preset in schema.DEPTHS.items():
        if key != 'matrix':
            assert not set(preset['sections']) & set(schema.MATRIX_ONLY)
    assert not set(schema.DEFAULT_SECTIONS) & set(schema.MATRIX_ONLY)


@pytest.mark.skipif(not (shutil.which('soffice') and shutil.which('pdftotext')),
                    reason='needs LibreOffice and poppler to render')
def test_a_long_head_to_head_renders_clean(tmp_path):
    import deckcheck
    card = _long_card()
    ratings = ['strong', 'partial', 'none', 'unknown']
    for index, row in enumerate(card['comparison']):
        row['ia'] = ratings[index % 4]
        row['competitor'] = ratings[(index // 4) % 4]
    card['meta']['depth'] = 'matrix'
    path = service.build(card, str(tmp_path))['path']
    _, faults = deckcheck.check(path)
    assert not faults, faults[:5]
