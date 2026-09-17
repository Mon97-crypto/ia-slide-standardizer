"""Tests for keeping and serving the original documents.

The library used to keep only the text extracted from an upload, which made
every entry searchable and none of them openable. These cover the document
being kept, handed back correctly, and handed back safely: this origin holds
the session cookie, so what it will render inline is a security decision.
"""
from __future__ import annotations

import io
import json
import os
import sys
import tempfile

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from ciq import files                        # noqa: E402

def _pptx() -> bytes:
    """A real deck. Extraction rejects anything that is not one, so a stand-in
    would only prove the upload path rejects rubbish."""
    from pptx import Presentation
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Blue Yonder head to head"
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def _pdf() -> bytes:
    from reportlab.pdfgen import canvas
    buffer = io.BytesIO()
    page = canvas.Canvas(buffer)
    page.drawString(72, 720, "Quarterly competitor brief")
    page.save()
    return buffer.getvalue()


PPTX = _pptx()
PDF = _pdf()


@pytest.fixture()
def client(monkeypatch):
    handle, path = tempfile.mkstemp(suffix=".db")
    os.close(handle)
    os.unlink(path)
    monkeypatch.setenv("CIQ_DB_PATH", path)
    monkeypatch.setenv("CIQ_SECRET_KEY", "test-secret")
    for module in [m for m in list(sys.modules)
                   if m == "ciq" or m.startswith("ciq.")] + ["app"]:
        sys.modules.pop(module, None)
    import app as application
    application.app.config["TESTING"] = True
    with application.app.test_client() as test_client:
        yield test_client
    for suffix in ("", "-wal", "-shm"):
        try:
            os.unlink(path + suffix)
        except OSError:
            pass


def upload(client, data=b"Competitor pricing notes.", name="notes.txt",
           **fields):
    payload = {"competitor": "o9", "title": "T", "category": "battlecard",
               **fields}
    if data is not None:
        payload["file"] = (io.BytesIO(data), name)
    response = client.post("/api/entries", data=payload,
                           content_type="multipart/form-data")
    assert response.status_code == 201, response.get_data(as_text=True)
    return response.get_json()["entry"]


# ─── naming and types ──────────────────────────────────────────────────────

@pytest.mark.parametrize("name,mime", [
    ("deck.pptx", "application/vnd.openxmlformats-officedocument."
                  "presentationml.presentation"),
    ("report.PDF", "application/pdf"),
    ("sheet.xlsx", "application/vnd.openxmlformats-officedocument."
                   "spreadsheetml.sheet"),
    ("notes.md", "text/plain"),
    ("mystery.bin", "application/octet-stream"),
    ("", "application/octet-stream"),
])
def test_the_type_comes_from_the_extension(name, mime):
    assert files.mime_for(name) == mime


@pytest.mark.parametrize("name", ["page.html", "logo.svg", "x.xhtml", "a.htm"])
def test_markup_is_never_rendered_in_place(name):
    """Inline on this origin, script in an uploaded file would run with the
    signed-in user's session. These download instead."""
    assert not files.opens_in_browser(files.mime_for(name))


@pytest.mark.parametrize("raw,expected", [
    ("../../etc/passwd", "passwd"),
    ("C:\\\\Users\\\\me\\\\deck.pptx", "deck.pptx"),
    ('quote"name.txt', "quotename.txt"),
    ("line\nbreak.txt", "linebreak.txt"),
    ("   ", "document"),
    ("", "document"),
])
def test_a_filename_cannot_escape_its_header_or_its_folder(raw, expected):
    assert files.safe_name(raw) == expected


# ─── keeping the document ──────────────────────────────────────────────────

def test_an_uploaded_file_can_be_opened_again(client):
    entry = upload(client, PPTX, "Q3 Deck.pptx")
    assert entry["has_file"] is True
    response = client.get(f"/api/entries/{entry['id']}/file")
    assert response.status_code == 200
    assert response.data == PPTX            # byte for byte, not the text
    assert "Q3 Deck.pptx" in response.headers["Content-Disposition"]
    assert response.headers["X-Content-Type-Options"] == "nosniff"


def test_a_pdf_opens_in_the_browser_and_a_deck_downloads(client):
    reading = upload(client, PDF, "brief.pdf")
    deck = upload(client, PPTX, "deck.pptx")
    assert "inline" in client.get(
        f"/api/entries/{reading['id']}/file").headers["Content-Disposition"]
    assert "attachment" in client.get(
        f"/api/entries/{deck['id']}/file").headers["Content-Disposition"]


def test_markup_stored_as_a_document_still_downloads(client):
    """The check has to hold through the route, not only in the helper. HTML
    cannot be uploaded, but it can be attached, so it can reach the store."""
    entry = upload(client, None, note="x")
    client.post(f"/api/entries/{entry['id']}/file",
                data={"file": (io.BytesIO(b"<script>alert(1)</script>"),
                               "notes.html")},
                content_type="multipart/form-data")
    response = client.get(f"/api/entries/{entry['id']}/file")
    assert "attachment" in response.headers["Content-Disposition"]
    assert response.headers["X-Content-Type-Options"] == "nosniff"
    assert "text/html" not in response.headers["Content-Type"]


def test_every_listing_says_which_entries_hold_a_file(client):
    with_file = upload(client, PPTX, "deck.pptx")
    note_only = upload(client, None, note="Just a note.")
    listed = {e["id"]: e for e in client.get("/api/entries").get_json()["entries"]}
    assert listed[with_file["id"]]["has_file"] is True
    assert listed[note_only["id"]]["has_file"] is False
    found = {e["id"]: e for e in
             client.get("/api/search?q=o9").get_json()["results"]}
    assert found[with_file["id"]]["has_file"] is True


def test_deleting_an_entry_takes_its_file_with_it(client):
    entry = upload(client, PPTX, "deck.pptx")
    client.delete(f"/api/entries/{entry['id']}")
    assert client.get(f"/api/entries/{entry['id']}/file").status_code == 404


def test_clearing_the_library_leaves_no_files_behind(client):
    import app as application
    from ciq import db
    entry = upload(client, PPTX, "deck.pptx")
    db.clear_all(application.store())
    assert db.get_file(application.store(), entry["id"]) is None


# ─── entries whose file was never kept ─────────────────────────────────────

def test_the_extracted_text_opens_when_the_file_is_not_held(client):
    """Everything added before files were kept, and every note, still has to
    be openable, or the library has entries nobody can read."""
    entry = upload(client, None, note="Deal notes from the Q3 review.")
    assert entry["has_file"] is False
    response = client.get(f"/api/entries/{entry['id']}/text")
    assert response.status_code == 200
    assert response.mimetype == "text/plain"
    # One charset, not two: passing a full type to send_file made Werkzeug
    # append a second one and the header came out malformed.
    assert response.headers["Content-Type"].count("charset") == 1
    assert response.headers["X-Content-Type-Options"] == "nosniff"
    body = response.get_data(as_text=True)
    assert "Deal notes from the Q3 review." in body
    assert "T" in body                       # the title, as a heading


def test_an_entry_without_an_original_still_serves_a_file(client):
    """Every entry opens. Where the document is gone, what comes back is a PDF
    built from the text the library does hold, not a dead link."""
    import pypdf
    entry = upload(client, None, note="Seven figure licences, per node.")
    response = client.get(f"/api/entries/{entry['id']}/file")
    assert response.status_code == 200
    assert response.mimetype == "application/pdf"
    assert response.data[:5] == b"%PDF-"
    text = pypdf.PdfReader(io.BytesIO(response.data)).pages[0].extract_text()
    assert "Seven figure licences" in text
    assert "T" in text                              # the entry title


def test_the_generated_file_says_it_is_not_the_original(client):
    """It gets forwarded to people who never saw the entry, so the page has to
    carry the caveat rather than the card it came from."""
    import pypdf
    entry = upload(client, None, note="A note.")
    response = client.get(f"/api/entries/{entry['id']}/file")
    text = pypdf.PdfReader(io.BytesIO(response.data)).pages[0].extract_text()
    assert "note written into the library" in text


def test_an_entry_whose_upload_predates_storage_explains_the_gap(client):
    import pypdf
    from ciq import db
    import app as application
    entry = upload(client, b"Extracted deck text about markdowns.", "old.txt")
    db.delete_file(application.store(), entry["id"])      # as an old row looks
    response = client.get(f"/api/entries/{entry['id']}/file")
    text = pypdf.PdfReader(io.BytesIO(response.data)).pages[0].extract_text()
    assert "old.txt" in text and "not held" in text
    assert "Extracted deck text about markdowns." in text


def test_an_entry_with_nothing_in_it_still_renders(client):
    """A link that failed to fetch leaves a title and not much else. It must
    still produce a file rather than a 500."""
    import pypdf
    entry = upload(client, None, note="x")
    from ciq import db
    import app as application
    db.update_entry(application.store(), entry["id"], {"content": "", "note": ""})
    response = client.get(f"/api/entries/{entry['id']}/file")
    assert response.status_code == 200
    text = pypdf.PdfReader(io.BytesIO(response.data)).pages[0].extract_text()
    assert "No text was extracted" in text


# ─── downloading ───────────────────────────────────────────────────────────

def test_every_entry_can_be_downloaded(client):
    """Both kinds: the stored original and the generated stand-in."""
    stored = upload(client, PPTX, "deck.pptx")
    generated = upload(client, None, note="Just a note.")
    for entry in (stored, generated):
        response = client.get(f"/api/entries/{entry['id']}/file?download=1")
        assert response.status_code == 200
        assert "attachment" in response.headers["Content-Disposition"]
        assert response.data


def test_download_forces_a_save_on_a_file_that_would_open_in_place(client):
    """A PDF opens in the tab on Open file. Download has to override that or
    the two actions do the same thing."""
    entry = upload(client, PDF, "brief.pdf")
    opened = client.get(f"/api/entries/{entry['id']}/file")
    saved = client.get(f"/api/entries/{entry['id']}/file?download=1")
    assert "inline" in opened.headers["Content-Disposition"]
    assert "attachment" in saved.headers["Content-Disposition"]
    assert opened.data == saved.data == PDF


def test_a_downloaded_file_keeps_its_own_name(client):
    stored = upload(client, PPTX, "Q3 Deck.pptx")
    generated = upload(client, None, note="x", title="Pricing note")
    assert "Q3 Deck.pptx" in client.get(
        f"/api/entries/{stored['id']}/file?download=1"
    ).headers["Content-Disposition"]
    assert "Pricing-note.pdf" in client.get(
        f"/api/entries/{generated['id']}/file?download=1"
    ).headers["Content-Disposition"]


def test_a_missing_entry_is_a_404_not_a_crash(client):
    for path in ("file", "text"):
        assert client.get(f"/api/entries/nope/{path}").status_code == 404


# ─── putting a document back ───────────────────────────────────────────────

def test_a_file_can_be_attached_to_an_older_entry(client):
    """The fix for a library that predates this: the entry keeps its id, its
    analysis and its place, and gains the document."""
    entry = upload(client, None, note="Only a note was saved.")
    response = client.post(f"/api/entries/{entry['id']}/file",
                           data={"file": (io.BytesIO(PPTX), "recovered.pptx")},
                           content_type="multipart/form-data")
    assert response.status_code == 200
    assert response.get_json()["entry"]["has_file"] is True
    served = client.get(f"/api/entries/{entry['id']}/file")
    assert served.data == PPTX
    assert "recovered.pptx" in served.headers["Content-Disposition"]


def test_attaching_indexes_a_file_that_had_no_text(client):
    entry = upload(client, None, note="x")
    client.post(f"/api/entries/{entry['id']}/file",
                data={"file": (io.BytesIO(b"Blue Yonder markdown pricing tiers."),
                               "pricing.txt")},
                content_type="multipart/form-data")
    found = client.get("/api/search?q=markdown pricing tiers").get_json()
    assert entry["id"] in [r["id"] for r in found["results"]]


def test_attaching_does_not_rewrite_text_that_came_from_a_document(client):
    """Replacing indexed text would quietly change what searches, analyses and
    battlecards were built on."""
    entry = upload(client, b"The original extracted content.", "original.txt")
    client.post(f"/api/entries/{entry['id']}/file",
                data={"file": (io.BytesIO(b"Something else entirely."),
                               "replacement.txt")},
                content_type="multipart/form-data")
    after = client.get(f"/api/entries/{entry['id']}").get_json()["entry"]
    assert after["content"] == "The original extracted content."


@pytest.mark.parametrize("data,reason", [
    (None, "Choose a file"),
    (b"", "empty"),
])
def test_attaching_nothing_is_refused(client, data, reason):
    entry = upload(client, None, note="x")
    payload = {} if data is None else {"file": (io.BytesIO(data), "empty.txt")}
    response = client.post(f"/api/entries/{entry['id']}/file", data=payload,
                           content_type="multipart/form-data")
    assert response.status_code == 400
    assert reason.lower() in response.get_json()["error"].lower()


# ─── the backup, which is what makes "forever" the team's to hold ──────────

def _backup(client) -> bytes:
    response = client.get("/api/library/export")
    assert response.status_code == 200
    assert response.mimetype == "application/zip"
    return response.data


def test_the_backup_carries_the_documents_not_just_the_rows(client):
    """A JSON export of entry rows is the obvious thing to reach for as a
    backup, and it would have held none of the uploads."""
    import zipfile
    entry = upload(client, PPTX, "Q3 Deck.pptx")
    archive = zipfile.ZipFile(io.BytesIO(_backup(client)))
    names = archive.namelist()
    assert "library.json" in names
    stored = [n for n in names if n.startswith("files/")]
    assert len(stored) == 1
    assert archive.read(stored[0]) == PPTX
    manifest = json.loads(archive.read("library.json"))
    record = next(r for r in manifest["entries"] if r["id"] == entry["id"])
    assert record["file"]["name"] == "Q3 Deck.pptx"
    assert record["file"]["bytes"] == len(PPTX)


def test_a_library_survives_a_full_round_trip(client):
    """The whole point: wipe everything, restore the backup, and the documents
    are still openable byte for byte."""
    upload(client, PPTX, "deck.pptx", title="The deck")
    upload(client, None, title="The note", note="Seven figure licences.")
    saved = _backup(client)

    client.post("/api/library/clear")
    assert client.get("/api/entries").get_json()["entries"] == []

    response = client.post("/api/library/import",
                           data={"file": (io.BytesIO(saved), "backup.zip")},
                           content_type="multipart/form-data")
    body = response.get_json()
    assert body["imported"] == 2 and body["restored"] == 1

    entries = {e["title"]: e for e in
               client.get("/api/entries").get_json()["entries"]}
    assert entries["The note"]["has_file"] is False
    assert entries["The deck"]["has_file"] is True
    served = client.get(f"/api/entries/{entries['The deck']['id']}/file")
    assert served.data == PPTX
    assert "deck.pptx" in served.headers["Content-Disposition"]
    # The note still opens too, as the PDF built from its text.
    assert client.get(
        f"/api/entries/{entries['The note']['id']}/file").data[:5] == b"%PDF-"


def test_an_attached_file_is_in_the_next_backup(client):
    """Attach original has to be covered too, or the recovery path loses
    exactly the files someone went to the trouble of putting back."""
    import zipfile
    entry = upload(client, None, note="Only a note was saved.")
    client.post(f"/api/entries/{entry['id']}/file",
                data={"file": (io.BytesIO(PPTX), "recovered.pptx")},
                content_type="multipart/form-data")
    archive = zipfile.ZipFile(io.BytesIO(_backup(client)))
    stored = [n for n in archive.namelist() if n.startswith("files/")]
    assert len(stored) == 1 and archive.read(stored[0]) == PPTX


def test_an_older_json_export_still_imports(client):
    """Anyone holding one of the old exports must not be stranded by this."""
    payload = json.dumps([{"competitor": "o9", "title": "Old", "note": "n",
                           "content": "old text"}]).encode()
    response = client.post("/api/library/import",
                           data={"file": (io.BytesIO(payload), "old.json")},
                           content_type="multipart/form-data")
    body = response.get_json()
    assert body["imported"] == 1 and body["restored"] == 0
    assert "holds no documents" in body["status"]


@pytest.mark.parametrize("data,name,reason", [
    (b"not a zip at all", "x.json", "not valid JSON"),
    (b"PK\x03\x04garbage", "x.zip", "not a readable zip"),
])
def test_a_broken_backup_says_what_is_wrong(client, data, name, reason):
    response = client.post("/api/library/import",
                           data={"file": (io.BytesIO(data), name)},
                           content_type="multipart/form-data")
    assert response.status_code == 400
    assert reason in response.get_json()["error"]


def test_a_zip_without_a_manifest_is_refused(client):
    import zipfile
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("files/whatever", b"data")
    response = client.post(
        "/api/library/import",
        data={"file": (io.BytesIO(buffer.getvalue()), "x.zip")},
        content_type="multipart/form-data")
    assert response.status_code == 400
    assert "not a library backup" in response.get_json()["error"]


def test_a_backup_cannot_be_made_to_read_outside_itself(client):
    """The manifest is data from a file someone was handed. A path in it is
    never followed - only names the archive actually lists are read."""
    import zipfile
    from ciq import backup as backup_module
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("library.json", json.dumps({"entries": [
            {"title": "Evil", "content": "x",
             "file": {"path": "../../../../etc/passwd", "name": "p"}}]}))
    parsed = backup_module.read(buffer.getvalue())
    assert parsed.file_for(parsed.records[0]) is None


def test_the_download_is_named_for_the_day_it_was_taken(client):
    upload(client, PPTX, "deck.pptx")
    disposition = client.get(
        "/api/library/export").headers["Content-Disposition"]
    assert "ia-competitor-library-" in disposition and ".zip" in disposition


# ─── saying whether it will actually last ──────────────────────────────────

def test_saving_a_file_says_where_it_went(client, monkeypatch):
    """An ephemeral deployment looks exactly like a permanent one until a
    deploy erases it, so the moment of saving is where this belongs."""
    import ciq.config
    entry = upload(client, None, note="x")

    monkeypatch.setattr(ciq.config.Config, "storage_info",
                        classmethod(lambda cls: {"durable": True}))
    body = client.post(f"/api/entries/{entry['id']}/file",
                       data={"file": (io.BytesIO(PPTX), "a.pptx")},
                       content_type="multipart/form-data").get_json()
    assert body["durable"] is True
    assert "survives deploys" in body["status"]

    monkeypatch.setattr(ciq.config.Config, "storage_info",
                        classmethod(lambda cls: {"durable": False}))
    body = client.post(f"/api/entries/{entry['id']}/file",
                       data={"file": (io.BytesIO(PPTX), "b.pptx")},
                       content_type="multipart/form-data").get_json()
    assert body["durable"] is False
    assert "lost on the next deploy" in body["status"]


# ─── proving to the user that storage is real ──────────────────────────────

def test_storage_reports_what_the_database_actually_holds(client):
    """"Are my uploads being stored" deserves a count read out of the
    database, not a reassurance."""
    before = client.get("/api/storage").get_json()
    assert before["documents"] == 0 and before["documents_supported"] is True
    upload(client, PPTX, "deck.pptx")
    after = client.get("/api/storage").get_json()
    assert after["documents"] == 1
    assert after["document_bytes"] == len(PPTX)
    assert after["largest"][0]["file_name"] == "deck.pptx"
    assert after["backend"] in ("sqlite", "postgres")


def test_storage_survives_a_database_that_predates_documents(client):
    """A build older than the files table must report that plainly instead of
    failing: it is the difference between a bug and a deploy that never
    happened. On Postgres a query against a missing table also aborts the
    transaction, so the check has to come first."""
    import app as application
    from ciq import db
    conn = application.store()
    conn.execute("DROP TABLE IF EXISTS files")
    conn.commit()
    assert db.has_files_table(conn) is False
    assert db.file_stats(conn) == {"supported": False, "documents": 0,
                                   "document_bytes": 0}
    assert db.largest_files(conn) == []
    # The connection is still usable afterwards, not left in an aborted state.
    assert db.stats(conn)["entries"] == 0
